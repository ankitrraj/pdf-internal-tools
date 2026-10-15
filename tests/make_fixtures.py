"""Build the sample files the test suite runs against."""
import os, sys, zipfile, json, csv
from pathlib import Path
import pymupdf as fitz
from PIL import Image, ImageDraw

HERE = Path(__file__).resolve().parent
default = Path(os.environ.get("PDFBUILDER_TEST_DIR", HERE / "_work")) / "fixtures"
out = Path(sys.argv[1]) if len(sys.argv) > 1 else default
out.mkdir(parents=True, exist_ok=True)

# ---- a 6-page text PDF with bookmarks, a table and an image
doc = fitz.open()
for i in range(6):
    page = doc.new_page(width=595, height=842)
    page.insert_text((60, 80), f"Chapter {i+1}", fontsize=22, fontname="hebo")
    body = (f"This is page {i+1} of the sample document. It mentions "
            f"Confidential material and the word apple several times. "
            f"apple apple.")
    page.insert_textbox(fitz.Rect(60, 110, 535, 400), body * 3, fontsize=11,
                        fontname="helv")
    # a ruled table so table detection has something to find
    y = 430
    for r in range(4):
        for c in range(3):
            rect = fitz.Rect(60 + c*150, y + r*24, 60 + (c+1)*150, y + (r+1)*24)
            page.draw_rect(rect, color=(0.4,0.4,0.4), width=0.6)
            page.insert_text((rect.x0+5, rect.y1-7), f"r{r}c{c}", fontsize=9)
    page.insert_text((60, 700), "email: nobody@example.com  phone 555-0100",
                     fontsize=10)
img = Image.new("RGB", (400, 260), (235, 240, 250))
d = ImageDraw.Draw(img)
d.ellipse((40, 40, 360, 220), fill=(90, 140, 230))
d.text((150, 120), "PICTURE", fill="white")
p_img = out / "photo.png"; img.save(p_img)
doc[2].insert_image(fitz.Rect(300, 560, 520, 690), filename=str(p_img))
doc.set_toc([[1, "Chapter 1", 1], [2, "Section 1.1", 2], [1, "Chapter 2", 3],
             [1, "Chapter 3", 5]])
doc.set_metadata({"title": "Sample Document", "author": "Test Author",
                  "subject": "Testing", "keywords": "sample, test"})
doc.save(str(out / "sample.pdf"))
doc.close()

# ---- a 2-page PDF for merge/compare/insert
doc = fitz.open()
for i in range(2):
    page = doc.new_page(width=595, height=842)
    page.insert_text((60, 80), f"Second file, page {i+1}", fontsize=18)
    page.insert_textbox(fitz.Rect(60,110,535,300), "Different wording here. "*20,
                        fontsize=11)
doc.save(str(out / "other.pdf")); doc.close()

# ---- a near-copy of sample.pdf for the comparison tools
doc = fitz.open(str(out / "sample.pdf"))
doc[1].insert_text((60, 760), "THIS LINE ONLY EXISTS IN THE SECOND VERSION",
                   fontsize=13, color=(1,0,0))
doc.save(str(out / "sample-v2.pdf")); doc.close()

# ---- a PDF with form fields
doc = fitz.open(); page = doc.new_page(width=595, height=842)
page.insert_text((60, 70), "Application form", fontsize=20, fontname="hebo")
def widget(kind, name, rect, **kw):
    w = fitz.Widget(); w.rect = fitz.Rect(*rect); w.field_name = name
    w.field_type = kind; w.text_fontsize = 10
    for k, v in kw.items(): setattr(w, k, v)
    page.add_widget(w)
widget(fitz.PDF_WIDGET_TYPE_TEXT, "full_name", (60,110,350,132), field_value="")
widget(fitz.PDF_WIDGET_TYPE_TEXT, "date_of_birth", (60,150,250,172), field_value="")
widget(fitz.PDF_WIDGET_TYPE_CHECKBOX, "agree", (60,190,76,206), field_value=False)
w = fitz.Widget(); w.rect = fitz.Rect(60,230,250,252)
w.field_name = "country"; w.field_type = fitz.PDF_WIDGET_TYPE_COMBOBOX
w.choice_values = ["India","United Kingdom","United States"]
w.field_value = "India"; w.text_fontsize = 10; page.add_widget(w)
doc.save(str(out / "form.pdf")); doc.close()

# ---- a "scanned" PDF: page image, no text layer, slightly rotated
big = Image.new("RGB", (1240, 1754), "white")
d = ImageDraw.Draw(big)
d.text((120, 160), "SCANNED INVOICE", fill="black")
for i in range(14):
    d.text((120, 240 + i*42), f"Line item {i+1} ..... 100.00", fill=(20,20,20))
d.rectangle((100,120,1140,820), outline=(120,120,120), width=2)
big = big.rotate(-2.4, fillcolor="white", resample=Image.BICUBIC)
scan = out / "scan-page.png"; big.save(scan)
doc = fitz.open(); page = doc.new_page(width=595, height=842)
page.insert_image(page.rect, filename=str(scan))
blank = doc.new_page(width=595, height=842)   # a truly blank page
doc.save(str(out / "scanned.pdf")); doc.close()

# ---- an encrypted PDF (password: secret)
doc = fitz.open(str(out / "sample.pdf"))
doc.save(str(out / "locked.pdf"), encryption=fitz.PDF_ENCRYPT_AES_256,
         owner_pw="owner", user_pw="secret")
doc.close()

# ---- a damaged PDF
raw = (out / "sample.pdf").read_bytes()
(out / "broken.pdf").write_bytes(raw[:int(len(raw)*0.82)])

# ---- images in several formats
for name, mode, fmt in (("pic1.jpg","RGB","JPEG"), ("pic2.png","RGBA","PNG"),
                        ("pic3.webp","RGB","WEBP"), ("pic4.tif","RGB","TIFF"),
                        ("pic5.bmp","RGB","BMP")):
    im = Image.new(mode, (900, 640), (250, 245, 230) if mode=="RGB" else (250,245,230,255))
    dd = ImageDraw.Draw(im)
    dd.rectangle((60,60,840,580), outline=(40,60,120), width=6)
    dd.text((120, 300), name, fill=(20,40,90))
    im.save(out / name, fmt)

# ---- text-ish inputs
(out / "notes.txt").write_text("Plain text file\n" + "line of text\n"*40, encoding="utf-8")
(out / "doc.md").write_text("# Title\n\nSome **bold** text.\n\n"
    "| a | b |\n|---|---|\n| 1 | 2 |\n\n```\ncode block\n```\n", encoding="utf-8")
(out / "page.html").write_text("<html><body><h1>Hello</h1><p>A paragraph with "
    "<b>bold</b>.</p><table><tr><td>x</td><td>y</td></tr></table></body></html>",
    encoding="utf-8")
(out / "data.csv").write_text("name,qty\nwidget,3\ngadget,5\n", encoding="utf-8")
(out / "fields.json").write_text(json.dumps(
    {"full_name":"Ada Lovelace","date_of_birth":"10/12/1815","agree":"yes",
     "country":"United Kingdom"}), encoding="utf-8")
with open(out / "rows.csv","w",newline="",encoding="utf-8") as fh:
    w = csv.DictWriter(fh, fieldnames=["full_name","date_of_birth","agree","country"])
    w.writeheader()
    w.writerow({"full_name":"Grace Hopper","date_of_birth":"09/12/1906",
                "agree":"yes","country":"United States"})
    w.writerow({"full_name":"Alan Turing","date_of_birth":"23/06/1912",
                "agree":"no","country":"United Kingdom"})

# ---- an archive
with zipfile.ZipFile(out / "bundle.zip", "w") as z:
    z.write(out / "pic1.jpg", "pic1.jpg")
    z.write(out / "pic2.png", "pic2.png")
    z.write(out / "notes.txt", "notes.txt")

# ---- office documents, built with the pure-python writers
try:
    import docx
    d = docx.Document(); d.add_heading("Word heading", 0)
    d.add_paragraph("A paragraph of text in a Word document.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0,0).text="a"; t.cell(0,1).text="b"; t.cell(1,0).text="1"; t.cell(1,1).text="2"
    d.save(str(out / "letter.docx"))
except Exception as e: print("docx fixture skipped:", e)
try:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Numbers"
    ws.append(["item","count"]); ws.append(["alpha", 4]); ws.append(["beta", 9])
    wb.save(str(out / "sheet.xlsx"))
except Exception as e: print("xlsx fixture skipped:", e)
try:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Slide title"
    s.placeholders[1].text = "Bullet one\nBullet two"
    prs.save(str(out / "deck.pptx"))
except Exception as e: print("pptx fixture skipped:", e)

(out / "note.rtf").write_text(r"{\rtf1\ansi An RTF document\par second line\par}",
                              encoding="ascii")
print("fixtures:", len(list(out.iterdir())), "files")
