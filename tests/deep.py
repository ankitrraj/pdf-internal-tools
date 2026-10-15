import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))
WORK = Path(os.environ.get("PDFBUILDER_TEST_DIR", HERE / "_work"))
WORK.mkdir(parents=True, exist_ok=True)
import shutil
import zipfile
from pathlib import Path

FIX = WORK/"fixtures"; OUT = WORK/"deep"; shutil.rmtree(OUT, ignore_errors=True)
from pdfbuilder import registry; registry.load_all()
from pdfbuilder.registry import REGISTRY
from pdfbuilder.job import Ctx, ToolError
import pymupdf as fitz
from PIL import Image, ImageDraw

def run(tid, inputs, **over):
    t = REGISTRY[tid]
    ctx = Ctx(inputs, {**t.defaults(), **over}, OUT/tid, log=lambda m: None)
    try:
        return [Path(p) for p in (t.fn(ctx) or ctx.produced)], ctx
    finally:
        ctx.cleanup()

checks=[]
def ck(name, cond, detail=""):
    checks.append((name, bool(cond), detail))
    print(f"{'  ok  ' if cond else ' FAIL '} {name}   {detail}")

from pdfbuilder.core.ocr import deskew_image
base = Image.new("RGB",(1100,760),"white"); d=ImageDraw.Draw(base)
for i in range(16): d.line((90,90+i*40,1010,90+i*40), fill=(15,15,15), width=5)
for a in (-7.5,-6.0,-2.5,3.0,7.5):
    fixed,_ = deskew_image(base.rotate(a, fillcolor="white", resample=Image.BICUBIC))
    _, residual = deskew_image(fixed)
    ck(f"deskew straightens a {a:+.1f}° tilt", abs(residual) < 0.6,
       f"residual {residual:+.2f}°")

made, ctx = run("remove_blank", [FIX/"scanned.pdf"], action="just report")
ck("blank detector flags the empty page", "Page 2" in str(getattr(ctx,"summary",{})),
   str(getattr(ctx,"summary",{})))
for order in ("across then down","down then across"):
    made,_ = run("nup",[FIX/"sample.pdf"], per="4", order=order)
    doc=fitz.open(made[0]); txt=doc[0].get_text()
    ck(f"4-up [{order}]", doc.page_count==2 and all(f"Chapter {i}" in txt for i in range(1,5)))
    doc.close()
for behind in (False,True):
    made,_=run("watermark_text",[FIX/"sample.pdf"],behind=behind,
               text="BEHIND" if behind else "FRONT")
    doc=fitz.open(made[0])
    ck(f"watermark behind={behind}", ("BEHIND" if behind else "FRONT") in doc[0].get_text())
    doc.close()
made,_=run("split_page",[FIX/"sample.pdf"],direction="vertically (left | right)")
doc=fitz.open(made[0])
ck("split in half doubles the pages", doc.page_count==12, str(doc.page_count))
ck("left half keeps the heading", "Chapter 1" in doc[0].get_text()); doc.close()
made,_=run("toc_page",[FIX/"sample.pdf"],links=True)
doc=fitz.open(made[0]); links=doc[0].get_links()
ck("contents page has working links", len(links)>=3 and
   any(l.get("page",-1)>=1 for l in links), f"{len(links)} links"); doc.close()
made,_=run("bookmarks",[FIX/"sample.pdf"],
           outline="Part A @ 1\n  Chapter 1 @ 2\n    Detail @ 3\nPart B @ 5")
doc=fitz.open(made[0])
ck("bookmark nesting preserved", [e[0] for e in doc.get_toc()]==[1,2,3,1]); doc.close()
made,_=run("repair",[FIX/"broken.pdf"])
doc=fitz.open(made[0])
ck("repair recovers a truncated file", doc.page_count>=4 and "Chapter" in doc[0].get_text(),
   f"{doc.page_count} pages"); doc.close()
for kind in ("text","checkbox","radio button","dropdown","date","signature"):
    made,_=run("form_create_quick",[FIX/"sample.pdf"],kind=kind,
               name=f"f_{kind.split()[0]}",options="A\nB")
    doc=fitz.open(made[0]); names=[w.field_name for w in (doc[0].widgets() or [])]
    ck(f"create a {kind} field", any(n.startswith("f_") for n in names), str(names))
    doc.close()
made,_=run("restrictions",[FIX/"sample.pdf"],owner_pw="own",allow_copy=False,
           allow_print=False,allow_modify=False)
doc=fitz.open(made[0]); p=doc.permissions
ck("copying blocked", not p & fitz.PDF_PERM_COPY)
ck("printing blocked", not p & fitz.PDF_PERM_PRINT)
ck("accessibility still allowed", p & fitz.PDF_PERM_ACCESSIBILITY); doc.close()
made,ctx=run("compare_report",[FIX/"sample.pdf",FIX/"sample-v2.pdf"],dpi=80)
doc=fitz.open(made[0]); ck("comparison sheet per page", doc.page_count==6); doc.close()
ck("text diff names the added line",
   "THIS LINE ONLY EXISTS" in [m for m in made if m.suffix==".txt"][0].read_text())
made,_=run("search_text",[FIX/"sample.pdf"],query=r"\d{3}-\d{4}",regex=True,save=True)
ck("regex search finds the phone number","555-0100" in made[0].read_text())
made,_=run("search_text",[FIX/"sample.pdf"],query="apple",whole=True,save=True)
ck("whole-word search returns hits", made[0].read_text().count("p1:")>0)
evil=OUT/"z"/"evil.zip"; evil.parent.mkdir(parents=True,exist_ok=True)
with zipfile.ZipFile(evil,"w") as z:
    z.writestr("../../escaped.txt","nope"); z.write(FIX/"pic1.jpg","ok.jpg")
made,_=run("zip_to_pdf",[evil])
ck("zip import blocks path traversal",
   not (OUT.parent/"escaped.txt").exists() and not Path("/tmp/escaped.txt").exists())
ck("zip import converts the safe entry", bool(made) and made[0].exists())
made,_=run("pdf_to_docx",[FIX/"sample.pdf"],pages="1-2")
import docx
ck("PDF → Word keeps the text",
   "Chapter 1" in "\n".join(p.text for p in docx.Document(str(made[0])).paragraphs))
made,_=run("pdf_to_xlsx",[FIX/"sample.pdf"],pages="1")
import openpyxl
wb=openpyxl.load_workbook(str(made[0]))
cells=[c.value for ws in wb for row in ws.iter_rows() for c in row if c.value]
ck("PDF → Excel extracts the table", any(str(c)=="r0c0" for c in cells), f"{len(cells)} cells")
made,_=run("pdf_to_csv",[FIX/"sample.pdf"],pages="1")
ck("PDF → CSV extracts rows","r0c0" in made[0].read_text())
made,_=run("protect",[FIX/"sample.pdf"],user_pw="pw123",encryption="AES-128")
doc=fitz.open(made[0]); ck("AES-128 encryption applied", doc.needs_pass); doc.close()
made,_=run("extract_images",[FIX/"sample.pdf"])
im=Image.open([m for m in made if m.suffix in (".png",".jpg",".jpeg")][0])
ck("embedded image extracted at full size", im.width>=400, f"{im.size}")
bad=[c for c in checks if not c[1]]
print(f"\n{len(checks)-len(bad)}/{len(checks)} deep behaviour checks passed")